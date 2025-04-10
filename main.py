import openai
from dotenv import load_dotenv
from moviepy import AudioFileClip
from docx import Document
from tqdm import tqdm
import os

# Configurar la API de OpenAI
load_dotenv()
openai.api_key = os.getenv("OPENAI_API_KEY")


# Función para dividir el audio en fragmentos usando moviepy
def split_audio(file_path, segment_duration=700):
    """
    Divide el audio en fragmentos usando moviepy.
    :param file_path: Ruta del archivo de audio.
    :param segment_duration: Duración de cada fragmento en segundos.
    :return: Lista de subclips de audio.
    """
    audio = AudioFileClip(file_path)
    duration = int(audio.duration)  # Duración total en segundos
    fragments = []

    for start in range(0, duration, segment_duration):
        end = min(start + segment_duration, duration)
        fragment = audio.subclipped(start, end)
        fragments.append(fragment)

    return fragments


# Función para transcribir un fragmento
def transcribe_audio_fragment(fragment, index):
    """
    Transcribe un fragmento de audio usando la API de OpenAI Whisper.
    :param fragment: Subclip de audio (moviepy).
    :param index: Índice del fragmento para el archivo temporal.
    :return: Texto transcrito.
    """
    # Guardar el fragmento como archivo temporal
    temp_file = f"temp_fragment_{index}.mp3"
    fragment.write_audiofile(temp_file, codec="libmp3lame")

    # Enviar a la API
    with open(temp_file, "rb") as audio_file:
        client = openai.OpenAI()
        response = client.audio.transcriptions.create(
            model="whisper-1", file=audio_file, prompt="The audio is a song"
        )

    # Eliminar el archivo temporal
    os.remove(temp_file)

    # Retornar la transcripción
    return response.text


# Función principal
def transcribe_audio_to_word(file_path, segment_duration=700):
    """
    Transcribe un archivo de audio largo dividiéndolo en fragmentos y guarda el texto en un archivo Word.
    :param file_path: Ruta del archivo de audio.
    :param segment_duration: Duración de los fragmentos en segundos.
    """
    # Validar que el archivo existe
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"El archivo de audio '{file_path}' no existe.")

    # Obtener el nombre del archivo de audio
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    output_docx = os.path.join("output", f"{base_name}.docx")

    # Dividir el audio en fragmentos
    print("Dividiendo el audio en fragmentos...")
    fragments = split_audio(file_path, segment_duration)

    # Crear documento de Word
    document = Document()
    document.add_heading("Transcripción de Audio", level=1)

    # Procesar cada fragmento del audio con barra de progreso
    print("Procesando los fragmentos de audio...")
    for index, fragment in enumerate(tqdm(fragments, desc="Progreso")):
        transcription = transcribe_audio_fragment(fragment, index)
        document.add_paragraph(transcription)
        document.add_paragraph("\n---\n")

    # Guardar el documento Word
    os.makedirs("output", exist_ok=True)
    document.save(output_docx)
    print(f"Transcripción completa guardada en '{output_docx}'")


# Parámetros de entrada
name = input("Ingrese el nombre del audio a convertir:")
audio_file_path = f"input/{name}.mp3"  # Cambiar por la ruta real del archivo de audio

# Transcribir el audio y guardar en Word
transcribe_audio_to_word(audio_file_path, segment_duration=700)
